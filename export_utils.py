"""Export helpers for RewardAI Professional."""
from __future__ import annotations

from io import BytesIO
from typing import Any

import pandas as pd


def excel_bytes(results: dict[str, Any]) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        for sheet_name in ["analysis", "risk_flags", "promotions", "budget_summary", "grade_summary", "department_summary", "scenario_comparison"]:
            df = results.get(sheet_name)
            if isinstance(df, pd.DataFrame):
                safe_name = sheet_name[:31].replace("_", " ").title()[:31]
                df.to_excel(writer, sheet_name=safe_name, index=False)
        metrics_df = pd.DataFrame(list(results.get("metrics", {}).items()), columns=["Metric", "Value"])
        metrics_df.to_excel(writer, sheet_name="Executive Summary", index=False)
        workbook = writer.book
        header_fmt = workbook.add_format({"bold": True, "bg_color": "#153C3C", "font_color": "white", "border": 1})
        money_fmt = workbook.add_format({"num_format": "#,##0", "border": 1})
        pct_fmt = workbook.add_format({"num_format": "0.0%", "border": 1})
        for ws in writer.sheets.values():
            ws.freeze_panes(1, 0)
            ws.set_row(0, 22, header_fmt)
            ws.set_column(0, 0, 18)
            ws.set_column(1, 20, 16)
    return output.getvalue()


def csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8")


def docx_bytes(report_text: str) -> bytes | None:
    try:
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        doc.add_heading("RewardAI Professional Report", level=0)
        for line in report_text.splitlines():
            if line.startswith("# "):
                doc.add_heading(line.replace("# ", ""), level=1)
            elif line.startswith("## "):
                doc.add_heading(line.replace("## ", ""), level=2)
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line.strip():
                p = doc.add_paragraph(line.replace("**", ""))
                for run in p.runs:
                    run.font.size = Pt(10)
        out = BytesIO()
        doc.save(out)
        return out.getvalue()
    except Exception:
        return None
